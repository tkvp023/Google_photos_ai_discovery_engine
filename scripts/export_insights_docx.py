import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        # Parse table rows
        parsed_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            # Check if separator row (e.g. :---, ---)
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                continue
            parsed_rows.append(cells)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            # Pad short rows
            for r in parsed_rows:
                while len(r) < num_cols:
                    r.append('')
            
            tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True
            
            for row_idx, r in enumerate(parsed_rows):
                for col_idx, text in enumerate(r):
                    cell = tbl.cell(row_idx, col_idx)
                    # Clean markdown bold/italics
                    clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                    clean_text = re.sub(r'\*(.*?)\*', r'\1', clean_text)
                    clean_text = re.sub(r'`(.*?)`', r'\1', clean_text)
                    cell.text = clean_text
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if row_idx == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    
                    if row_idx == 0:
                        set_cell_background(cell, '1E293B')
                    elif row_idx % 2 == 1:
                        set_cell_background(cell, 'F8FAFC')

        table_rows = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.startswith('```'):
            if in_table:
                flush_table()
            in_code_block = not in_code_block
            i += 1
            continue
            
        if in_code_block:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Check table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Check headings
        if line.startswith('# '):
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:])
            run.font.color.rgb = RGBColor(15, 23, 42)
            run.bold = True
        elif line.startswith('## '):
            p = doc.add_heading(level=2)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:])
            run.font.color.rgb = RGBColor(30, 41, 59)
            run.bold = True
        elif line.startswith('### '):
            p = doc.add_heading(level=3)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[4:])
            run.font.color.rgb = RGBColor(51, 65, 85)
            run.bold = True
        elif line.startswith('#### '):
            p = doc.add_heading(level=4)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:])
            run.font.color.rgb = RGBColor(71, 85, 105)
            run.bold = True
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            quote_text = line[2:]
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', quote_text)
            clean_text = re.sub(r'`(.*?)`', r'\1', clean_text)
            run = p.add_run(clean_text)
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif line.startswith('---'):
            pass  # divider
        elif line.strip() == '':
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            # Add runs supporting simple bold
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    subparts = re.split(r'(`.*?`)', part)
                    for sp in subparts:
                        if sp.startswith('`') and sp.endswith('`'):
                            run = p.add_run(sp[1:-1])
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(225, 29, 72)
                        else:
                            run = p.add_run(sp)
                            run.font.size = Pt(10)

        i += 1

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully generated {docx_path}")

if __name__ == '__main__':
    convert_md_to_docx('insights_photos.md', 'insights_photos.docx')
